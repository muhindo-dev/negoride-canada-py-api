"""Admin — Payments & Finance (spec §19.1.7, §13).

Roles: finance, super_admin. Every endpoint writes audit_logs (admin views of
personal data are logged too, PIPEDA). All amounts are integer cents, CAD.
List endpoints accept ?from=YYYY-MM-DD&to=YYYY-MM-DD (UTC, `to` inclusive),
?page=&per_page= and ?format=csv | ?format=xlsx (CSV / Excel export of the whole filtered set).

    GET  /api/admin/finance/receipts                 list (+csv)  ?q=number|customer_id|driver_id|ride_type
    GET  /api/admin/finance/receipts/{id}            detail (+ credit notes, payment, refunds)
    GET  /api/admin/finance/receipts/{id}/pdf        PDF
    POST /api/admin/receipts/{id}/resend             re-email, SAME number (alias under /finance/)
    GET  /api/admin/finance/credit-notes             list (+csv)
    GET  /api/admin/finance/credit-notes/{id}/pdf    PDF
    GET  /api/admin/finance/payments                 ride payments list (+csv) ?capture_status=&purpose=&ride_type=
    GET  /api/admin/finance/payments/summary         authorizations / captures / releases / refunds totals
    GET  /api/admin/finance/refunds                  refunds + releases (+csv) ?kind=
    GET  /api/admin/finance/payouts                  payout requests overview + list (+csv) ?status=
    GET  /api/admin/finance/commission               commission revenue by day (+csv)
    GET  /api/admin/finance/tax                      tax collected per province, net of credit notes (+csv)
    GET  /api/admin/finance/statements               driver statements (+csv)
    GET  /api/admin/finance/statements/{id}/pdf      statement PDF
    POST /api/admin/finance/statements/run           {week_start?} build statements for a week (job)
    GET  /api/admin/finance/reconciliation           DB vs payment provider, per payment (+csv)
    GET  /api/admin/finance/tax-rates                effective-dated provincial rates (?province=)
    POST /api/admin/finance/tax-rates                {province, name, gst_bp, pst_bp, hst_bp, qst_bp, effective_from, effective_to?}
    PUT  /api/admin/finance/tax-rates/{id}           same fields (partial); audited, no overlaps per province
    GET  /api/admin/finance/tip-receipts             tip receipts NR-TIP-… (+csv/xlsx) ?q=&ride_type=&ride_id=&customer_id=&driver_id=
    GET  /api/admin/finance/tip-receipts/{id}/pdf    tip receipt PDF (audited)
"""
import csv
import io
from collections import defaultdict
from datetime import date, datetime, timedelta

from flask import Blueprint, Response, request

from backend import jobs
from backend.models import db
from backend.models.money import CreditNote, Receipt, Refund, RidePayment
from backend.models.payout_request import PayoutRequest
from backend.models.statements import DriverStatement
from backend.services import receipts as RC
from backend.services.audit import audit
from backend.services.payments.gateway import GatewayError, get_gateway
from backend.utils.auth import admin_role_required
from backend.utils.money import fmt, to_cents
from backend.utils.response import error_response, paginated_response, success_response

admin_finance_bp = Blueprint('admin_finance', __name__)
ROLES = ('finance',)   # super_admin always passes
CSV_MAX_ROWS = 50000


# ── helpers ─────────────────────────────────────────────────────────────────

def _page():
    page = max(1, request.args.get('page', 1, type=int))
    per = min(200, max(1, request.args.get('per_page', 25, type=int)))
    return page, per


def _range():
    """(start, end) naive UTC datetimes or None."""
    def parse(v):
        try:
            return datetime.combine(date.fromisoformat(v), datetime.min.time()) if v else None
        except ValueError:
            return None
    start = parse(request.args.get('from'))
    end = parse(request.args.get('to'))
    return start, (end + timedelta(days=1)) if end else None


def _dated(q, col):
    start, end = _range()
    if start:
        q = q.filter(col >= start)
    if end:
        q = q.filter(col < end)
    return q


def _export_format():
    f = (request.args.get('format') or '').lower()
    return f if f in ('csv', 'xlsx') else None


def _want_csv():
    """True for any file export (?format=csv or ?format=xlsx)."""
    return _export_format() is not None


XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _xlsx_bytes(sheet_name, columns, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = sheet_name[:31] or 'export'
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([_xlsx_value(r.get(c)) for c in columns])
    ws.freeze_panes = 'A2'
    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(10, min(40, len(col) + 4))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_value(v):
    if v is None:
        return None
    if isinstance(v, (int, float, str, bool, datetime, date)):
        # Excel formula injection: neutralise cells that start with = + - @
        if isinstance(v, str) and v[:1] in ('=', '+', '-', '@') and not _numeric(v):
            return "'" + v
        return v
    return str(v)


def _numeric(v):
    try:
        float(v)
        return True
    except ValueError:
        return False


def _csv(filename, columns, rows, admin, action):
    """File export of a filtered set: CSV (default) or Excel with ?format=xlsx."""
    fmt_ = _export_format() or 'csv'
    rows = list(rows)
    audit('finance.export', admin, 'finance', action,
          meta={'rows': len(rows), 'format': fmt_, 'filters': {k: v for k, v in request.args.items()}})
    db.session.commit()
    stamp = datetime.utcnow().strftime('%Y%m%d-%H%M')
    if fmt_ == 'xlsx':
        return Response(_xlsx_bytes(filename, columns, rows), mimetype=XLSX_MIME, headers={
            'Content-Disposition': f'attachment; filename="{filename}-{stamp}.xlsx"', 'Cache-Control': 'no-store'})
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(columns)
    for r in rows:
        w.writerow(['' if r.get(c) is None else r.get(c) for c in columns])
    return Response(buf.getvalue(), mimetype='text/csv', headers={
        'Content-Disposition': f'attachment; filename="{filename}-{stamp}.csv"', 'Cache-Control': 'no-store'})


def _viewed(admin, what, entity_id=None, personal=False):
    audit('personal_data.view' if personal else 'finance.view', admin, 'finance' if not personal else what,
          entity_id if entity_id is not None else what,
          meta={'view': what, 'filters': {k: v for k, v in request.args.items()}})
    db.session.commit()


def _list(q, order_col, to_row, admin, name, columns):
    if _want_csv():
        rows = [to_row(x) for x in q.order_by(order_col.desc()).limit(CSV_MAX_ROWS)]
        return _csv(name, columns, rows, admin, name)
    page, per = _page()
    total = q.count()
    rows = [to_row(x) for x in q.order_by(order_col.desc()).offset((page - 1) * per).limit(per)]
    _viewed(admin, name, personal=True)
    return paginated_response(rows, total, page, per, message=name.replace('_', ' ').title())


def _pdf(blob, filename):
    return Response(blob, mimetype='application/pdf', headers={
        'Content-Disposition': f'inline; filename="{filename}"', 'Cache-Control': 'private, no-store'})


# ── receipts ────────────────────────────────────────────────────────────────

RECEIPT_COLS = ['id', 'number', 'issued_at', 'ride_type', 'ride_id', 'customer_id', 'driver_id', 'province',
                'subtotal_cents', 'tax_cents', 'tip_cents', 'total_cents', 'credited_cents', 'net_total_cents',
                'payment_method', 'intent_id', 'emailed_at', 'email_count', 'voided_at']


def _credited_map(receipt_ids):
    out = defaultdict(int)
    if receipt_ids:
        for cn in CreditNote.query.filter(CreditNote.receipt_id.in_(receipt_ids)):
            out[cn.receipt_id] += int(cn.amount_cents)
    return out


def _receipt_row(r, credited=0):
    t = r.totals or {}
    return {'id': r.id, 'number': r.number, 'issued_at': RC._iso(r.issued_at), 'ride_type': r.ride_type,
            'ride_id': r.ride_id, 'customer_id': r.customer_id, 'driver_id': r.driver_id,
            'province': t.get('province'), 'subtotal_cents': t.get('subtotal_cents'), 'tax_cents': t.get('tax_cents'),
            'tip_cents': t.get('tip_cents'), 'total_cents': int(r.total_cents), 'credited_cents': credited,
            'net_total_cents': int(r.total_cents) - credited, 'payment_method': t.get('payment_method'),
            'intent_id': (t.get('internal') or {}).get('intent_id'), 'emailed_at': RC._iso(r.emailed_at),
            'email_count': r.email_count, 'voided_at': RC._iso(r.voided_at)}


@admin_finance_bp.route('/api/admin/finance/receipts', methods=['GET'])
@admin_role_required(*ROLES)
def receipts_list(admin):
    q = _dated(Receipt.query, Receipt.issued_at)
    a = request.args
    if a.get('q'):
        q = q.filter(Receipt.number.like(f"%{a['q'].strip()}%"))
    for f in ('customer_id', 'driver_id', 'ride_id'):
        if a.get(f, type=int):
            q = q.filter(getattr(Receipt, f) == a.get(f, type=int))
    if a.get('ride_type'):
        q = q.filter(Receipt.ride_type == a['ride_type'])
    if _want_csv():
        rows = q.order_by(Receipt.id.desc()).limit(CSV_MAX_ROWS).all()
        cm = _credited_map([r.id for r in rows])
        return _csv('receipts', RECEIPT_COLS, [_receipt_row(r, cm[r.id]) for r in rows], admin, 'receipts')
    page, per = _page()
    total = q.count()
    rows = q.order_by(Receipt.id.desc()).offset((page - 1) * per).limit(per).all()
    cm = _credited_map([r.id for r in rows])
    _viewed(admin, 'receipts', personal=True)
    return paginated_response([_receipt_row(r, cm[r.id]) for r in rows], total, page, per, message='Receipts')


@admin_finance_bp.route('/api/admin/finance/receipts/<int:receipt_id>', methods=['GET'])
@admin_role_required(*ROLES)
def receipt_detail(admin, receipt_id):
    r = db.session.get(Receipt, receipt_id)
    if r is None:
        return error_response('Receipt not found', status_code=404)
    data = RC.receipt_payload(r, viewer='admin', base_path=f'/api/admin/finance/receipts/{r.id}/pdf')
    data['totals'] = r.totals            # admins see the full snapshot (incl. internal)
    data['customer_id'], data['driver_id'] = r.customer_id, r.driver_id
    rp = db.session.get(RidePayment, r.ride_payment_id) if r.ride_payment_id else None
    data['payment'] = rp.to_dict() if rp else None
    data['refunds'] = [x.to_dict() for x in Refund.query.filter_by(ride_type=r.ride_type, ride_id=r.ride_id)
                       .order_by(Refund.id.asc())]
    audit('personal_data.view', admin, 'receipt', r.id, meta={'number': r.number})
    db.session.commit()
    return success_response('Receipt', data)


@admin_finance_bp.route('/api/admin/finance/receipts/<int:receipt_id>/pdf', methods=['GET'])
@admin_role_required(*ROLES)
def receipt_pdf(admin, receipt_id):
    r = db.session.get(Receipt, receipt_id)
    if r is None:
        return error_response('Receipt not found', status_code=404)
    blob = RC.ensure_pdf(r)
    audit('personal_data.view', admin, 'receipt', r.id, meta={'number': r.number, 'format': 'pdf'})
    db.session.commit()
    return _pdf(blob, f'NegoRide-receipt-{r.number}.pdf')


@admin_finance_bp.route('/api/admin/receipts/<int:receipt_id>/resend', methods=['POST'])
@admin_finance_bp.route('/api/admin/finance/receipts/<int:receipt_id>/resend', methods=['POST'])
@admin_role_required('finance', 'support', 'ops')
def receipt_resend(admin, receipt_id):
    r = db.session.get(Receipt, receipt_id)
    if r is None:
        return error_response('Receipt not found', status_code=404)
    # Sending (PDF + provider call) happens in a job; audited there (receipt.resend).
    jobs.enqueue('backend.services.receipts.resend_job', r.id, admin.id)
    db.session.rollback()
    r = db.session.get(Receipt, receipt_id)
    return success_response(f'Receipt {r.number} is being re-sent', {
        'id': r.id, 'number': r.number, 'queued': True, 'email_count': r.email_count,
        'emailed_at': RC._iso(r.emailed_at)}, status_code=202)


# ── tip receipts (NR-TIP-…) — admin list + PDF (additive, admin console) ────

TIP_RECEIPT_COLS = ['id', 'number', 'issued_at', 'ride_type', 'ride_id', 'receipt_id', 'ride_payment_id',
                    'customer_id', 'driver_id', 'amount_cents', 'currency', 'paid_at', 'payment_method', 'emailed_at']


def _tip_receipt_row(t):
    tot = t.totals or {}
    return {'id': t.id, 'number': t.number, 'issued_at': RC._iso(t.issued_at), 'ride_type': t.ride_type,
            'ride_id': t.ride_id, 'receipt_id': t.receipt_id, 'ride_payment_id': t.ride_payment_id,
            'customer_id': t.customer_id, 'driver_id': t.driver_id, 'amount_cents': int(t.amount_cents),
            'currency': t.currency, 'paid_at': tot.get('paid_at'), 'payment_method': tot.get('payment_method'),
            'emailed_at': RC._iso(t.emailed_at)}


@admin_finance_bp.route('/api/admin/finance/tip-receipts', methods=['GET'])
@admin_role_required(*ROLES)
def tip_receipts_list(admin):
    """Tip receipts (NR-TIP-YYYY-NNNNNN) ?q=number&ride_type=&ride_id=&customer_id=&driver_id= (+csv/xlsx)."""
    from backend.models.money import TipReceipt
    q = _dated(TipReceipt.query, TipReceipt.issued_at)
    a = request.args
    if a.get('q'):
        q = q.filter(TipReceipt.number.like(f"%{a['q'].strip()}%"))
    for f in ('customer_id', 'driver_id', 'ride_id'):
        if a.get(f, type=int):
            q = q.filter(getattr(TipReceipt, f) == a.get(f, type=int))
    if a.get('ride_type'):
        q = q.filter(TipReceipt.ride_type == a['ride_type'])
    return _list(q, TipReceipt.id, _tip_receipt_row, admin, 'tip_receipts', TIP_RECEIPT_COLS)


@admin_finance_bp.route('/api/admin/finance/tip-receipts/<int:tip_receipt_id>/pdf', methods=['GET'])
@admin_role_required(*ROLES)
def tip_receipt_pdf_admin(admin, tip_receipt_id):
    from backend.models.money import TipReceipt
    t = db.session.get(TipReceipt, tip_receipt_id)
    if t is None:
        return error_response('Tip receipt not found', data={'error_code': 'not_found'}, status_code=404)
    blob = RC.ensure_tip_pdf(t)
    audit('personal_data.view', admin, 'tip_receipt', t.id, meta={'number': t.number, 'format': 'pdf'})
    db.session.commit()
    return _pdf(blob, f'NegoRide-tip-{t.number}.pdf')


# ── tax rates (effective-dated, audited) ────────────────────────────────────

def _tax_row(t):
    total = int(t.gst_bp or 0) + int(t.pst_bp or 0) + int(t.hst_bp or 0) + int(t.qst_bp or 0)
    return {'id': t.id, 'province': t.province, 'name': t.name, 'gst_bp': t.gst_bp, 'pst_bp': t.pst_bp,
            'hst_bp': t.hst_bp, 'qst_bp': t.qst_bp, 'total_bp': total,
            'effective_from': t.effective_from.isoformat() if t.effective_from else None,
            'effective_to': t.effective_to.isoformat() if t.effective_to else None,
            'current': bool(t.effective_from and t.effective_from <= date.today()
                            and (t.effective_to is None or t.effective_to >= date.today()))}


def _tax_values(data, row=None):
    from backend.services.receipts import PROVINCES
    out = {}
    prov = (data.get('province') if 'province' in data else (row.province if row else '')) or ''
    prov = str(prov).strip().upper()
    if prov not in PROVINCES:
        raise ValueError('province must be a Canadian province/territory code (ON, QC, …).')
    out['province'] = prov
    out['name'] = str(data.get('name') or (row.name if row else '') or f'{prov} sales tax')[:60]
    for k in ('gst_bp', 'pst_bp', 'hst_bp', 'qst_bp'):
        v = data.get(k, getattr(row, k, 0) if row else 0)
        try:
            v = int(v or 0)
        except (TypeError, ValueError):
            raise ValueError(f'{k} must be an integer (basis points, 500 = 5 %).')
        if v < 0 or v > 5000:
            raise ValueError(f'{k} must be between 0 and 5000 basis points.')
        out[k] = v
    if out['hst_bp'] and (out['gst_bp'] or out['pst_bp']):
        raise ValueError('HST replaces GST + PST — set either hst_bp or gst_bp/pst_bp.')
    for k in ('effective_from', 'effective_to'):
        v = data.get(k, getattr(row, k, None) if row else None)
        if isinstance(v, str):
            v = v.strip() or None
            try:
                v = date.fromisoformat(v) if v else None
            except ValueError:
                raise ValueError(f'{k} must be YYYY-MM-DD.')
        out[k] = v
    if not out['effective_from']:
        raise ValueError('effective_from is required (YYYY-MM-DD).')
    if out['effective_to'] and out['effective_to'] < out['effective_from']:
        raise ValueError('effective_to must be on or after effective_from.')
    return out


def _overlaps(values, exclude_id=None):
    from backend.models.money import TaxRate
    q = TaxRate.query.filter(TaxRate.province == values['province'])
    if exclude_id:
        q = q.filter(TaxRate.id != exclude_id)
    far = date(9999, 12, 31)
    for t in q.all():
        a1, a2 = t.effective_from, t.effective_to or far
        b1, b2 = values['effective_from'], values['effective_to'] or far
        if a1 <= b2 and b1 <= a2:
            return t
    return None


@admin_finance_bp.route('/api/admin/finance/tax-rates', methods=['GET'])
@admin_role_required(*ROLES)
def tax_rates_list(admin):
    from backend.models.money import TaxRate
    q = TaxRate.query
    if request.args.get('province'):
        q = q.filter(TaxRate.province == request.args['province'].upper())
    rows = [_tax_row(t) for t in q.order_by(TaxRate.province.asc(), TaxRate.effective_from.desc())]
    _viewed(admin, 'tax_rates')
    return success_response('Tax rates', {'items': rows})


@admin_finance_bp.route('/api/admin/finance/tax-rates', methods=['POST'])
@admin_role_required(*ROLES)
def tax_rates_create(admin):
    """New effective-dated rate. To change a rate from a date, close the current
    row (PUT effective_to = day before) and create the new one; rows of one
    province may not overlap."""
    from backend.models.money import TaxRate
    data = request.get_json(silent=True) or {}
    try:
        values = _tax_values(data)
    except ValueError as e:
        return error_response(str(e), data={'error_code': 'bad_tax_rate'})
    clash = _overlaps(values)
    if clash is not None:
        return error_response(f'Overlaps rate #{clash.id} ({clash.effective_from} → {clash.effective_to or "open"}). '
                              'Close it first (effective_to).', data={'error_code': 'overlap', 'rate_id': clash.id},
                              status_code=409)
    t = TaxRate(**values)
    db.session.add(t)
    db.session.flush()
    audit('finance.tax_rate_create', admin, 'tax_rate', t.id, after=_tax_row(t))
    db.session.commit()
    return success_response('Tax rate created', _tax_row(t), status_code=201)


@admin_finance_bp.route('/api/admin/finance/tax-rates/<int:rate_id>', methods=['PUT', 'PATCH'])
@admin_role_required(*ROLES)
def tax_rates_update(admin, rate_id):
    from backend.models.money import TaxRate
    t = db.session.get(TaxRate, rate_id)
    if t is None:
        return error_response('Tax rate not found', status_code=404)
    data = request.get_json(silent=True) or {}
    try:
        values = _tax_values(data, row=t)
    except ValueError as e:
        return error_response(str(e), data={'error_code': 'bad_tax_rate'})
    clash = _overlaps(values, exclude_id=t.id)
    if clash is not None:
        return error_response(f'Overlaps rate #{clash.id}.', data={'error_code': 'overlap', 'rate_id': clash.id},
                              status_code=409)
    before = _tax_row(t)
    for k, v in values.items():
        setattr(t, k, v)
    db.session.flush()
    audit('finance.tax_rate_update', admin, 'tax_rate', t.id, before=before, after=_tax_row(t))
    db.session.commit()
    return success_response('Tax rate updated', _tax_row(t))


# ── credit notes ────────────────────────────────────────────────────────────

CN_COLS = ['id', 'number', 'issued_at', 'receipt_id', 'receipt_number', 'refund_id', 'customer_id', 'province',
           'amount_cents', 'subtotal_cents', 'tax_cents', 'reason', 'emailed_at']


def _cn_row(c):
    t = c.totals or {}
    return {'id': c.id, 'number': c.number, 'issued_at': RC._iso(c.issued_at), 'receipt_id': c.receipt_id,
            'receipt_number': t.get('receipt_number'), 'refund_id': c.refund_id, 'customer_id': c.customer_id,
            'province': t.get('province'), 'amount_cents': int(c.amount_cents), 'subtotal_cents': t.get('subtotal_cents'),
            'tax_cents': int(c.tax_cents or 0), 'reason': c.reason, 'emailed_at': RC._iso(c.emailed_at)}


@admin_finance_bp.route('/api/admin/finance/credit-notes', methods=['GET'])
@admin_role_required(*ROLES)
def credit_notes_list(admin):
    q = _dated(CreditNote.query, CreditNote.issued_at)
    if request.args.get('receipt_id', type=int):
        q = q.filter(CreditNote.receipt_id == request.args.get('receipt_id', type=int))
    return _list(q, CreditNote.id, _cn_row, admin, 'credit_notes', CN_COLS)


@admin_finance_bp.route('/api/admin/finance/credit-notes/<int:cn_id>/pdf', methods=['GET'])
@admin_role_required(*ROLES)
def credit_note_pdf(admin, cn_id):
    c = db.session.get(CreditNote, cn_id)
    if c is None:
        return error_response('Credit note not found', status_code=404)
    blob = RC.ensure_credit_note_pdf(c)
    audit('personal_data.view', admin, 'credit_note', c.id, meta={'number': c.number, 'format': 'pdf'})
    db.session.commit()
    return _pdf(blob, f'NegoRide-credit-note-{c.number}.pdf')


# ── payments ────────────────────────────────────────────────────────────────

PAY_COLS = ['id', 'created_at', 'ride_type', 'ride_id', 'purpose', 'customer_id', 'driver_id', 'capture_method',
            'capture_status', 'fare_cents', 'fees_cents', 'amount_authorized_cents', 'amount_captured_cents',
            'amount_refunded_cents', 'net_cents', 'payment_method', 'intent_id', 'authorized_at', 'captured_at']


def _pay_row(p):
    return {'id': p.id, 'created_at': RC._iso(p.created_at), 'ride_type': p.ride_type, 'ride_id': p.ride_id,
            'purpose': p.purpose, 'customer_id': p.customer_id, 'driver_id': p.driver_id,
            'capture_method': p.capture_method, 'capture_status': p.capture_status, 'fare_cents': p.fare_cents,
            'fees_cents': p.fees_cents, 'amount_authorized_cents': p.amount_authorized_cents,
            'amount_captured_cents': p.amount_captured_cents, 'amount_refunded_cents': p.amount_refunded_cents,
            'net_cents': p.net_cents, 'payment_method': RC.payment_method_label(p), 'intent_id': p.intent_id,
            'authorized_at': RC._iso(p.authorized_at), 'captured_at': RC._iso(p.captured_at)}


def _payments_query():
    q = _dated(RidePayment.query, RidePayment.created_at)
    for f in ('capture_status', 'purpose', 'ride_type'):
        if request.args.get(f):
            q = q.filter(getattr(RidePayment, f) == request.args[f])
    return q


@admin_finance_bp.route('/api/admin/finance/payments', methods=['GET'])
@admin_role_required(*ROLES)
def payments_list(admin):
    return _list(_payments_query(), RidePayment.id, _pay_row, admin, 'payments', PAY_COLS)


@admin_finance_bp.route('/api/admin/finance/payments/summary', methods=['GET'])
@admin_role_required(*ROLES)
def payments_summary(admin):
    by_status = defaultdict(lambda: {'count': 0, 'authorized_cents': 0, 'captured_cents': 0, 'refunded_cents': 0})
    by_purpose = defaultdict(lambda: {'count': 0, 'captured_cents': 0})
    totals = {'authorized_cents': 0, 'captured_cents': 0, 'refunded_cents': 0, 'open_holds_cents': 0,
              'open_holds_count': 0, 'expiring_24h_count': 0}
    soon = datetime.utcnow() + timedelta(hours=24)
    for p in _payments_query().all():
        s = by_status[p.capture_status]
        s['count'] += 1
        s['authorized_cents'] += int(p.amount_authorized_cents or 0)
        s['captured_cents'] += int(p.amount_captured_cents or 0)
        s['refunded_cents'] += int(p.amount_refunded_cents or 0)
        by_purpose[p.purpose]['count'] += 1
        by_purpose[p.purpose]['captured_cents'] += int(p.amount_captured_cents or 0)
        totals['authorized_cents'] += int(p.amount_authorized_cents or 0)
        totals['captured_cents'] += int(p.amount_captured_cents or 0)
        totals['refunded_cents'] += int(p.amount_refunded_cents or 0)
        if p.capture_status == 'authorized':
            totals['open_holds_cents'] += int(p.amount_authorized_cents or 0)
            totals['open_holds_count'] += 1
            if p.auth_expires_at and p.auth_expires_at <= soon:
                totals['expiring_24h_count'] += 1
    rq = _dated(Refund.query, Refund.created_at)
    releases = sum(int(r.amount_cents) for r in rq.filter(Refund.kind == 'release', Refund.status == 'succeeded'))
    totals['released_cents'] = releases
    totals['net_collected_cents'] = totals['captured_cents'] - totals['refunded_cents']
    if _want_csv():
        cols = ['capture_status', 'count', 'authorized_cents', 'captured_cents', 'refunded_cents']
        return _csv('payments-summary', cols, [{'capture_status': k, **v} for k, v in sorted(by_status.items())],
                    admin, 'payments_summary')
    _viewed(admin, 'payments_summary')
    return success_response('Payments summary', {'totals': totals, 'by_status': dict(by_status),
                                                 'by_purpose': dict(by_purpose)})


# ── refunds ─────────────────────────────────────────────────────────────────

REFUND_COLS = ['id', 'created_at', 'kind', 'status', 'ride_type', 'ride_id', 'ride_payment_id', 'customer_id',
               'amount_cents', 'rule_id', 'reason', 'initiated_by_type', 'initiated_by', 'provider_refund_id',
               'credit_note_id']


def _refund_row(r):
    d = r.to_dict()
    return {c: d.get(c) for c in REFUND_COLS}


@admin_finance_bp.route('/api/admin/finance/refunds', methods=['GET'])
@admin_role_required(*ROLES)
def refunds_list(admin):
    q = _dated(Refund.query, Refund.created_at)
    if request.args.get('kind'):
        q = q.filter(Refund.kind == request.args['kind'])
    return _list(q, Refund.id, _refund_row, admin, 'refunds', REFUND_COLS)


# ── payouts ─────────────────────────────────────────────────────────────────

PAYOUT_COLS = ['id', 'requested_at', 'user_id', 'status', 'payout_method', 'amount_cents', 'fee_cents',
               'net_cents', 'currency', 'stripe_transfer_id', 'stripe_payout_id', 'processed_at', 'failure_reason']


def _payout_row(p):
    return {'id': p.id, 'requested_at': RC._iso(p.requested_at), 'user_id': p.user_id, 'status': p.status,
            'payout_method': p.payout_method, 'amount_cents': to_cents(p.amount), 'fee_cents': to_cents(p.fee_amount),
            'net_cents': to_cents(p.net_amount), 'currency': p.currency, 'stripe_transfer_id': p.stripe_transfer_id,
            'stripe_payout_id': p.stripe_payout_id, 'processed_at': RC._iso(p.processed_at),
            'failure_reason': p.failure_reason}


@admin_finance_bp.route('/api/admin/finance/payouts', methods=['GET'])
@admin_role_required(*ROLES)
def payouts(admin):
    q = _dated(PayoutRequest.query.filter(PayoutRequest.deleted_at.is_(None)), PayoutRequest.requested_at)
    if request.args.get('status'):
        q = q.filter(PayoutRequest.status == request.args['status'])
    if _want_csv():
        return _list(q, PayoutRequest.id, _payout_row, admin, 'payouts', PAYOUT_COLS)
    overview = defaultdict(lambda: {'count': 0, 'amount_cents': 0})
    for p in q.all():
        overview[p.status]['count'] += 1
        overview[p.status]['amount_cents'] += to_cents(p.amount)
    page, per = _page()
    total = q.count()
    rows = [_payout_row(p) for p in q.order_by(PayoutRequest.id.desc()).offset((page - 1) * per).limit(per)]
    _viewed(admin, 'payouts', personal=True)
    import math
    return success_response('Payouts', {'overview': dict(overview), 'current_page': page, 'data': rows,
                                        'per_page': per, 'total': total,
                                        'last_page': math.ceil(total / per) if per else 1})


# ── commission & tax ────────────────────────────────────────────────────────

@admin_finance_bp.route('/api/admin/finance/commission', methods=['GET'])
@admin_role_required(*ROLES)
def commission(admin):
    """Commission revenue from receipted rides (+ booking fees), by day."""
    days = defaultdict(lambda: {'rides': 0, 'fares_cents': 0, 'commission_cents': 0, 'booking_fees_cents': 0})
    for r in _dated(Receipt.query.filter(Receipt.voided_at.is_(None)), Receipt.issued_at):
        t = r.totals or {}
        internal = t.get('internal') or {}
        d = days[r.issued_at.strftime('%Y-%m-%d')]
        d['rides'] += 1
        d['fares_cents'] += int(internal.get('driver_fare_cents') or 0)
        d['commission_cents'] += int(internal.get('commission_cents') or 0)
        d['booking_fees_cents'] += sum(int(l.get('amount_incl_tax_cents') if l.get('amount_incl_tax_cents') is not None
                                           else l.get('amount_cents') or 0)
                                       for l in t.get('lines') or [] if l.get('code') == 'booking_fee')
    rows = [{'date': k, **v, 'revenue_cents': v['commission_cents'] + v['booking_fees_cents']}
            for k, v in sorted(days.items())]
    cols = ['date', 'rides', 'fares_cents', 'commission_cents', 'booking_fees_cents', 'revenue_cents']
    if _want_csv():
        return _csv('commission', cols, rows, admin, 'commission')
    totals = {c: sum(r[c] for r in rows) for c in cols[1:]}
    _viewed(admin, 'commission')
    return success_response('Commission revenue', {'days': rows, 'totals': totals})


@admin_finance_bp.route('/api/admin/finance/tax', methods=['GET'])
@admin_role_required(*ROLES)
def tax_collected(admin):
    """Tax collected per province and component, net of credit notes."""
    prov = defaultdict(lambda: {'receipts': 0, 'taxable_cents': 0, 'tax_cents': 0, 'credited_tax_cents': 0,
                                'components': defaultdict(int)})
    for r in _dated(Receipt.query.filter(Receipt.voided_at.is_(None)), Receipt.issued_at):
        t = r.totals or {}
        p = prov[t.get('province') or '??']
        p['receipts'] += 1
        p['taxable_cents'] += int(t.get('subtotal_cents') or 0)
        p['tax_cents'] += int(t.get('tax_cents') or 0)
        for x in t.get('taxes') or []:
            p['components'][x['code']] += int(x['amount_cents'])
    for c in _dated(CreditNote.query, CreditNote.issued_at):
        t = c.totals or {}
        p = prov[t.get('province') or '??']
        p['credited_tax_cents'] += int(c.tax_cents or 0)
        for x in t.get('taxes') or []:
            p['components'][x['code']] -= int(x['amount_cents'])
    rows = []
    for code, v in sorted(prov.items()):
        comps = dict(v['components'])
        rows.append({'province': code, 'province_name': RC.PROVINCES.get(code, code), 'receipts': v['receipts'],
                     'taxable_cents': v['taxable_cents'], 'tax_cents': v['tax_cents'],
                     'credited_tax_cents': v['credited_tax_cents'], 'net_tax_cents': v['tax_cents'] - v['credited_tax_cents'],
                     'gst_cents': comps.get('GST', 0), 'hst_cents': comps.get('HST', 0),
                     'pst_cents': comps.get('PST', 0), 'qst_cents': comps.get('QST', 0)})
    cols = ['province', 'province_name', 'receipts', 'taxable_cents', 'tax_cents', 'credited_tax_cents',
            'net_tax_cents', 'gst_cents', 'hst_cents', 'pst_cents', 'qst_cents']
    if _want_csv():
        return _csv('tax-by-province', cols, rows, admin, 'tax')
    _viewed(admin, 'tax')
    return success_response('Tax collected', {'provinces': rows,
                                              'totals': {c: sum(r[c] for r in rows) for c in cols[2:]}})


# ── driver statements ───────────────────────────────────────────────────────

ST_COLS = ['id', 'number', 'driver_id', 'period_start', 'gross_cents', 'commission_cents', 'fees_cents',
           'net_cents', 'payouts_cents', 'emailed_at']


def _st_row(s):
    d = s.to_dict(exclude=('totals',))
    return {c: d.get(c) for c in ST_COLS}


@admin_finance_bp.route('/api/admin/finance/statements', methods=['GET'])
@admin_role_required(*ROLES)
def statements(admin):
    q = DriverStatement.query
    if request.args.get('driver_id', type=int):
        q = q.filter(DriverStatement.driver_id == request.args.get('driver_id', type=int))
    start, end = _range()
    if start:
        q = q.filter(DriverStatement.period_start >= start.date())
    if end:
        q = q.filter(DriverStatement.period_start < end.date())
    return _list(q, DriverStatement.id, _st_row, admin, 'statements', ST_COLS)


@admin_finance_bp.route('/api/admin/finance/statements/<int:st_id>/pdf', methods=['GET'])
@admin_role_required(*ROLES)
def statement_pdf(admin, st_id):
    from backend.services import receipt_jobs
    s = db.session.get(DriverStatement, st_id)
    if s is None:
        return error_response('Statement not found', status_code=404)
    blob = receipt_jobs.ensure_statement_pdf(s)
    audit('personal_data.view', admin, 'driver_statement', s.id, meta={'number': s.number})
    db.session.commit()
    return _pdf(blob, f'NegoRide-statement-{s.number}.pdf')


@admin_finance_bp.route('/api/admin/finance/statements/run', methods=['POST'])
@admin_role_required(*ROLES)
def statements_run(admin):
    data = request.get_json(silent=True) or {}
    week = data.get('week_start')
    if week:
        try:
            date.fromisoformat(week)
        except ValueError:
            return error_response('week_start must be YYYY-MM-DD')
    audit('finance.statements_run', admin, 'driver_statement', week or 'last_week', meta={'week_start': week})
    db.session.commit()
    jobs.enqueue('backend.services.receipt_jobs.weekly_driver_statements', week)
    return success_response('Statement run queued', {'week_start': week})


# ── reconciliation ──────────────────────────────────────────────────────────

RECON_COLS = ['ride_payment_id', 'intent_id', 'purpose', 'ride_type', 'ride_id', 'capture_status',
              'db_captured_cents', 'provider_captured_cents', 'db_refunded_cents', 'provider_refunded_cents',
              'receipt_number', 'receipt_charged_cents', 'status', 'issues']


def _reconcile(p, gw):
    row = {'ride_payment_id': p.id, 'intent_id': p.intent_id, 'purpose': p.purpose, 'ride_type': p.ride_type,
           'ride_id': p.ride_id, 'capture_status': p.capture_status,
           'db_captured_cents': int(p.amount_captured_cents or 0), 'db_refunded_cents': int(p.amount_refunded_cents or 0),
           'provider_captured_cents': None, 'provider_refunded_cents': None, 'receipt_number': None,
           'receipt_charged_cents': None}
    issues = []
    try:
        prov = gw.retrieve_payment_totals(p.intent_id)
        row['provider_captured_cents'] = prov['amount_received']
        row['provider_refunded_cents'] = prov['amount_refunded']
        if prov['amount_received'] != row['db_captured_cents']:
            issues.append('captured_mismatch')
        if prov['amount_refunded'] != row['db_refunded_cents']:
            issues.append('refunded_mismatch')
        if (prov.get('currency') or 'cad').lower() != 'cad':
            issues.append('currency_mismatch')
    except (GatewayError, AttributeError) as e:
        issues.append(f'provider_error: {e}'[:200])
    if p.purpose == 'ride':
        rc = Receipt.query.filter_by(ride_type=p.ride_type, ride_id=p.ride_id).first()
        if rc is not None:
            row['receipt_number'] = rc.number
            charged = int((rc.totals or {}).get('ride_total_cents') or 0)
            row['receipt_charged_cents'] = charged
            if (rc.totals or {}).get('internal', {}).get('ride_payment_id') == p.id and charged != row['db_captured_cents']:
                issues.append('receipt_mismatch')
    row['issues'] = ';'.join(issues)
    row['status'] = 'ok' if not issues else ('error' if any(i.startswith('provider_error') for i in issues) else 'mismatch')
    return row


@admin_finance_bp.route('/api/admin/finance/reconciliation', methods=['GET'])
@admin_role_required(*ROLES)
def reconciliation(admin):
    """Compare our captured/refunded totals with the payment provider, one
    payment at a time (bounded page: provider calls are made synchronously)."""
    from backend.models.money import CAPTURED_STATES
    q = _dated(RidePayment.query.filter(RidePayment.intent_id.isnot(None), RidePayment.provider != 'offline',
                                        RidePayment.capture_status.in_(CAPTURED_STATES)),
               RidePayment.captured_at)
    if request.args.get('purpose'):
        q = q.filter(RidePayment.purpose == request.args['purpose'])
    page = max(1, request.args.get('page', 1, type=int))
    per = min(100, max(1, request.args.get('per_page', 50, type=int)))
    total = q.count()
    gw = get_gateway()
    rows = [_reconcile(p, gw) for p in q.order_by(RidePayment.id.desc()).offset((page - 1) * per).limit(per)]
    only = request.args.get('only')
    if only == 'issues':
        rows = [r for r in rows if r['status'] != 'ok']
    if _want_csv():
        return _csv('stripe-reconciliation', RECON_COLS, rows, admin, 'reconciliation')
    summary = {'checked': len(rows), 'ok': sum(r['status'] == 'ok' for r in rows),
               'mismatch': sum(r['status'] == 'mismatch' for r in rows),
               'error': sum(r['status'] == 'error' for r in rows),
               'db_captured_cents': sum(r['db_captured_cents'] for r in rows),
               'provider_captured_cents': sum(r['provider_captured_cents'] or 0 for r in rows),
               'db_refunded_cents': sum(r['db_refunded_cents'] for r in rows),
               'provider_refunded_cents': sum(r['provider_refunded_cents'] or 0 for r in rows),
               'provider': getattr(gw, 'name', 'unknown')}
    audit('finance.reconciliation', admin, 'finance', 'reconciliation', meta={'page': page, 'summary': summary})
    db.session.commit()
    import math
    return success_response('Reconciliation', {'summary': summary, 'current_page': page, 'data': rows,
                                               'per_page': per, 'total': total,
                                               'last_page': math.ceil(total / per) if per else 1})

"""Audit gaps — trip effects copy, notification escalation / admins / webhooks /
bounces / Live Activities / template overrides, receipts (tips, driver copy,
logo, sweeper, admin issue, tax rates, xlsx), province lookup, service config."""
import base64
import io
import json
import re
from datetime import date, datetime, timedelta

import pytest

from backend import jobs
from backend.models import db
from backend.models.money import Receipt, RidePayment, TaxRate, TipReceipt
from backend.models.negotiation import Negotiation
from backend.models.notification import LiveActivityToken, Notification, NotificationDelivery
from backend.models.platform import AuditLog, WebhookEvent
from backend.models.trip_booking import TripBooking
from backend.models.user import AdminUser
from backend.services import receipts as RC
from backend.services import receipt_jobs
from backend.services import settings_service as S
from backend.services import trip_effects as TE
from backend.services.notify import dispatcher, email_provider
from backend.services.notify import email_status
from backend.services.notify import live_activity as LA
from backend.services.payments import payment_service as PS
from backend.services.payments.gateway import get_gateway
from backend.utils.province import province_at
from tests.conftest import body
from tests.test_carhire_flow import PICKUP, _confirmed_ride, create_negotiation
from tests.test_receipts import paid_carhire


def _send(event, uid, ctx=None):
    ids = dispatcher.notify_now(event, [uid], ctx or {})
    db.session.rollback()
    return db.session.get(Notification, ids[0]) if ids else None


def _latest(uid, key):
    return Notification.query.filter_by(user_id=uid, event_key=key).order_by(Notification.id.desc()).first()


# ── province lookup (unit) ─────────────────────────────────────────────────

@pytest.mark.parametrize('city,point,code', [
    ('Vancouver', (49.2827, -123.1207), 'BC'), ('Montréal', (45.5019, -73.5674), 'QC'),
    ('Calgary', (51.0447, -114.0719), 'AB'), ('Halifax', (44.6488, -63.5752), 'NS'),
    ('Toronto', (43.6532, -79.3832), 'ON'), ('Winnipeg', (49.8951, -97.1384), 'MB'),
    ('Regina', (50.4452, -104.6189), 'SK'), ('Ottawa', (45.4215, -75.6972), 'ON'),
    ('Gatineau', (45.4765, -75.7013), 'QC'), ('Moncton', (46.0878, -64.7782), 'NB'),
    ('Charlottetown', (46.2382, -63.1311), 'PE'), ("St. John's", (47.5615, -52.7126), 'NL'),
    ('Whitehorse', (60.7212, -135.0568), 'YT'), ('Kenora', (49.767, -94.49), 'ON'),
    ('New York', (40.71, -74.0), None), ('Seattle', (47.6, -122.3), None),
])
def test_province_at(city, point, code):
    assert province_at(*point) == code, city


def test_pickup_province_is_stamped_at_creation(client, auth, make_user):
    customer, driver = make_user('customer'), make_user('driver')
    neg_id = create_negotiation(client, auth, customer, driver)
    assert db.session.get(Negotiation, neg_id).pickup_province == 'ON'
    r = client.post('/api/bookings', headers=auth(customer), json={
        'service_type': 'airport', 'pickup_lat': 45.4706, 'pickup_lng': -73.7408, 'pickup_address': 'YUL',
        'destination_lat': 45.5019, 'destination_lng': -73.5674, 'destination_address': 'Downtown',
        'customer_proposed_price': 5000,
        'scheduled_at': (datetime.utcnow() + timedelta(hours=3)).strftime('%Y-%m-%d %H:%M:%S')})
    from backend.models.scheduled_booking import ScheduledBooking
    assert db.session.get(ScheduledBooking, body(r)['data']['id']).pickup_province == 'QC'
    from tests.exp_helpers import make_trip
    from backend.services import rideshare_service as RS
    trip = make_trip(driver, start=(49.2827, -123.1207))
    b = RS.create_booking(customer, trip.id, pickup={'lat': 49.28, 'lng': -123.12, 'address': 'Waterfront'})
    assert b.pickup_province == 'BC'
    assert RC.province_of('rideshare_booking', b) == ('BC', 'ride')


# ── 2. trip effects copy ───────────────────────────────────────────────────

def test_completed_driver_copy_shows_net_earning(client, auth, make_user):
    customer, driver, neg_id, rp = _confirmed_ride(client, auth, make_user)
    client.post(f'/api/rides/carhire/{neg_id}/en-route', headers=auth(driver))
    client.post(f'/api/rides/carhire/{neg_id}/arrived', headers=auth(driver), json={'lat': PICKUP[0], 'lng': PICKUP[1]})
    client.post(f'/api/rides/carhire/{neg_id}/start', headers=auth(driver))
    assert body(client.post(f'/api/rides/carhire/{neg_id}/complete', headers=auth(driver)))['code'] == 1
    n = _latest(driver.id, 'ride.completed')
    assert '$19.80 added to your earnings' in n.body      # $22.00 − 10 %
    assert 'Rate your trip' in _latest(customer.id, 'ride.completed').body


def test_cancel_copy_says_refund_only_when_money_was_secured(client, auth, make_user):
    customer, driver = make_user('customer'), make_user('driver')
    neg_id = create_negotiation(client, auth, customer, driver)
    client.post(f'/api/rides/carhire/{neg_id}/cancel', headers=auth(driver))
    n = _latest(customer.id, 'ride.cancelled')
    assert n is not None and 'refund' not in n.body.lower()
    customer, driver, neg_id, rp = _confirmed_ride(client, auth, make_user)
    client.post(f'/api/rides/carhire/{neg_id}/cancel', headers=auth(driver))
    assert 'Full refund issued.' in _latest(customer.id, 'ride.cancelled').body


def test_rideshare_started_and_trip_completion_notify_with_local_times(client, auth, make_user):
    from tests.exp_helpers import make_trip
    from backend.services import rideshare_service as RS
    driver = make_user('driver')
    customer = make_user('customer', timezone='America/Vancouver')
    trip = make_trip(driver, depart_in=timedelta(hours=3))
    b = RS.create_booking(customer, trip.id, seats=1)
    r = client.post(f'/api/rides/rideshare_booking/{b.id}/pay', headers=auth(customer))
    rp = db.session.get(RidePayment, body(r)['data']['ride_payment_id'])
    PS.handle_stripe_event(get_gateway().simulate_customer_pays(rp.checkout_session_id))
    db.session.rollback()
    conf = _latest(customer.id, 'rideshare.booking_confirmed')
    expected = TE.local_dt(customer.id, trip.departure_at, TE.DEPARTURE_FMT)
    assert expected in conf.body and 'UTC' not in conf.body
    drv_conf = _latest(driver.id, 'rideshare.booking_confirmed')
    assert TE.local_dt(driver.id, trip.departure_at, TE.DEPARTURE_FMT) in drv_conf.body   # Toronto default
    assert client.post(f'/api/rides/rideshare_booking/{b.id}/start', headers=auth(driver)).get_json()['code'] == 1
    assert body(client.post(f'/api/rides/rideshare_trip/{trip.id}/start', headers=auth(driver)))['code'] == 1
    db.session.rollback()
    assert db.session.get(TripBooking, b.id).trip_stage == 'RIDING'
    assert _latest(customer.id, 'ride.started') is not None
    assert body(client.post(f'/api/rides/rideshare_trip/{trip.id}/complete', headers=auth(driver)))['code'] == 1
    n = _latest(driver.id, 'ride.completed')
    assert n is not None and '$22.50 added to your earnings' in n.body


# ── 3. notifications ────────────────────────────────────────────────────────

def test_every_critical_push_escalates_to_sms(make_user):
    u = make_user('customer')
    jobs.DEFERRED.clear()
    _send('ride.driver_arriving', u.id, {'eta_min': 1})       # critical, no sms_fallback flag
    assert any(p.endswith('escalate_if_unopened') for _, p, _, _ in jobs.DEFERRED)
    jobs.DEFERRED.clear()
    _send('ride.driver_en_route', u.id, {'driver_first': 'A'})   # not critical
    assert not any(p.endswith('escalate_if_unopened') for _, p, _, _ in jobs.DEFERRED)


def test_notify_admins_helper(make_user):
    from backend.services.notify import notify_admins
    from backend.services import realtime
    ops = make_user('admin', admin_roles='ops')
    fin = make_user('admin', admin_roles='finance')
    realtime.SENT.clear()
    ids = notify_admins('admin.background_check_review', {'name': 'Sam D.', 'result': 'consider', 'check_id': 7},
                        roles=('ops', 'safety_reviewer'))
    db.session.commit()
    assert ops.id in ids and fin.id not in ids
    n = _latest(ops.id, 'admin.background_check_review')
    assert n is not None and 'Sam D.: background check consider' in n.body
    assert any(e == 'admin.background_check_review' and room == 'admin:ops' for e, _, room in realtime.SENT)


def test_twilio_status_webhook_is_persisted_then_processed(client, make_user, monkeypatch):
    from backend.routes.verify import twilio_signature
    monkeypatch.setenv('TWILIO_AUTH_TOKEN', 'tw_test_token')
    u = make_user('customer')
    n = _send('ride.cancelled', u.id, {'cancelled_by': 'driver'})
    d = NotificationDelivery.query.filter_by(notification_id=n.id, channel='sms').first()
    d.provider_message_id = f'SM{n.id}test'
    d.status = 'sent'
    db.session.commit()
    params = {'MessageSid': d.provider_message_id, 'MessageStatus': 'delivered'}
    url = 'http://localhost/api/webhooks/twilio/status'
    r = client.post('/api/webhooks/twilio/status', data=params,
                    headers={'X-Twilio-Signature': twilio_signature(url, params, 'tw_test_token')})
    assert r.status_code == 204
    db.session.rollback()
    row = WebhookEvent.query.filter_by(provider='twilio_status', event_id=f'{d.provider_message_id}:delivered').one()
    assert row.status == 'processed'
    assert db.session.get(NotificationDelivery, d.id).status == 'delivered'
    from backend.services import platform_jobs
    assert platform_jobs.WEBHOOK_PROCESSORS['twilio_status'].endswith('process_twilio_status')
    assert platform_jobs.WEBHOOK_PROCESSORS['postmark'].endswith('process_postmark_event')
    db.session.delete(row)
    db.session.commit()


def test_postmark_hard_bounce_suppresses_email_until_cleared(client, make_user, monkeypatch):
    monkeypatch.setenv('POSTMARK_WEBHOOK_USER', 'pm')
    monkeypatch.setenv('POSTMARK_WEBHOOK_PASSWORD', 'secret')
    h = {'Authorization': 'Basic ' + base64.b64encode(b'pm:secret').decode()}
    customer, driver, neg, rp = paid_carhire(make_user)
    r = client.post('/api/webhooks/postmark', headers=h, json={
        'RecordType': 'Bounce', 'Type': 'HardBounce', 'ID': 9001 + neg.id, 'MessageID': f'pm-{neg.id}',
        'Email': customer.email, 'Inactive': True})
    assert r.status_code == 204
    db.session.rollback()
    u = db.session.get(AdminUser, customer.id)
    assert u.email_bounced_at is not None and 'HardBounce' in u.email_bounce_reason
    row = WebhookEvent.query.filter_by(provider='postmark', event_id=f'pm-{neg.id}:Bounce:{9001 + neg.id}').one()
    assert row.status == 'processed'
    # email channel skipped for notifications …
    n = _send('refund.issued', customer.id, {'amount': '$5.00'})
    assert not NotificationDelivery.query.filter_by(notification_id=n.id, channel='email').count()
    # … and for the receipt (issued, not emailed, logged as suppressed, no retry scheduled)
    email_provider.OUTBOX.clear()
    jobs.DEFERRED.clear()
    receipt = RC.issue_and_send('carhire', neg.id)
    assert receipt is not None and not [m for m in email_provider.OUTBOX if m['to'] == customer.email]
    assert not any(p.endswith('retry_email') for _, p, _, _ in jobs.DEFERRED)
    nr = _latest(customer.id, 'ride.receipt')
    dl = NotificationDelivery.query.filter_by(notification_id=nr.id, channel='email').one()
    assert dl.status == 'failed' and dl.error.startswith('suppressed')
    # identity hook: changing / verifying the email lifts it
    assert email_status.clear_email_bounce(u, commit=True) is True
    db.session.rollback()
    assert db.session.get(AdminUser, customer.id).email_bounced_at is None
    db.session.delete(row)
    db.session.commit()


def test_connect_payout_webhook_emits_payout_sent(client, sign_stripe, make_user):
    from backend.models.payout_account import PayoutAccount
    from tests.test_carhire_flow import deliver_webhook
    driver = make_user('driver')
    acct = f'acct_test_{driver.id}'
    db.session.add(PayoutAccount(user_id=driver.id, stripe_account_id=acct, status='active', payouts_enabled=True))
    db.session.commit()
    try:
        ev = {'id': f'evt_po_{driver.id}', 'type': 'payout.paid', 'account': acct,
              'data': {'object': {'id': f'po_{driver.id}', 'object': 'payout', 'amount': 14620, 'currency': 'cad'}}}
        assert deliver_webhook(client, sign_stripe, ev).status_code == 200
        assert deliver_webhook(client, sign_stripe, {**ev, 'id': ev['id'] + 'b'}).status_code == 200  # replay other id
        db.session.rollback()
        rows = Notification.query.filter_by(user_id=driver.id, event_key='payout.sent').all()
        assert len(rows) == 1 and '$146.20' in rows[0].body
    finally:
        db.session.rollback()
        PayoutAccount.query.filter_by(user_id=driver.id).delete()
        WebhookEvent.query.filter(WebhookEvent.event_id.like(f'evt_po_{driver.id}%')).delete(synchronize_session=False)
        db.session.commit()


# ── 4. Live Activities ─────────────────────────────────────────────────────

def test_live_activity_register_update_and_end(client, auth, make_user):
    customer, driver, neg_id, rp = _confirmed_ride(client, auth, make_user)
    outsider = make_user('customer')
    payload = {'ride_type': 'carhire', 'ride_id': neg_id, 'activity_id': f'la-{neg_id}', 'push_token': 'tok123'}
    assert client.post('/api/devices/live-activity', headers=auth(outsider), json=payload).status_code == 403
    r = client.post('/api/devices/live-activity', headers=auth(customer), json=payload)
    assert r.status_code == 201 and 'push_token' not in body(r)['data'], r.get_json()
    LA.LA_LOG.clear()
    client.post(f'/api/rides/carhire/{neg_id}/en-route', headers=auth(driver))
    sent = [b for a, b in LA.LA_LOG if a == f'la-{neg_id}']
    assert sent and sent[-1]['event'] == 'update' and sent[-1]['event_updates']['stage'] == 'DRIVER_EN_ROUTE'
    assert set(sent[-1]['event_updates']) >= {'stage', 'title', 'eta_min', 'driver_first', 'progress', 'pin'}
    n = _latest(customer.id, 'ride.driver_en_route')
    assert NotificationDelivery.query.filter_by(notification_id=n.id, channel='live_activity', status='sent').count() == 1
    # ETA refresh from services/eta.py
    assert LA.push_update('carhire', neg_id, {'seconds': 240}) == 1
    assert LA.LA_LOG[-1][1]['event_updates']['eta_min'] == 4
    # terminal → end
    client.post(f'/api/rides/carhire/{neg_id}/cancel', headers=auth(customer))
    assert LA.LA_LOG[-1][1]['event'] == 'end' and 'dismissal_date' in LA.LA_LOG[-1][1]
    db.session.rollback()
    assert LiveActivityToken.query.filter_by(activity_id=f'la-{neg_id}').one().status == 'ended'
    # a rider without an activity gets no live_activity delivery row
    other = make_user('customer')
    n2 = _send('ride.driver_en_route', other.id, {'ride_type': 'carhire', 'ride_id': neg_id})
    assert not NotificationDelivery.query.filter_by(notification_id=n2.id, channel='live_activity').count()


# ── 8. settings & template overrides ──────────────────────────────────────

def test_services_enabled_in_app_config(client):
    cfg = body(client.get('/api/app/config'))['data']
    assert cfg['services'] == S.enabled_services()
    assert cfg['settings']['services.enabled']


def test_template_override_is_used_and_previewable(client, auth, make_user):
    from backend.models.notification import NotificationTemplateOverride as O
    ops, u = make_user('admin', admin_roles='ops'), make_user('customer')
    key = 'rating.reminder'
    try:
        bad = client.put(f'/api/admin/notifications/templates/{key}', headers=auth(ops),
                         json={'lang': 'en', 'body': 'Hi {% if %}'})
        assert body(bad)['data']['error_code'] == 'bad_template'
        r = client.put(f'/api/admin/notifications/templates/{key}', headers=auth(ops),
                       json={'lang': 'en', 'body': 'Loved riding with {{ other_first }}? Tell us!'})
        assert body(r)['code'] == 1, r.get_json()
        n = _send(key, u.id, {'other_first': 'Amara'})
        assert n.body == 'Loved riding with Amara? Tell us!' and n.title == 'How was your trip?'   # title falls back
        pv = body(client.post(f'/api/admin/notifications/templates/{key}/preview', headers=auth(ops),
                              json={'lang': 'en', 'context': {'other_first': 'Zoé'}}))['data']
        assert pv['body'] == 'Loved riding with Zoé? Tell us!'
        lst = body(client.get('/api/admin/notifications/templates', headers=auth(ops)))['data']['items']
        assert next(i for i in lst if i['event_key'] == key)['overrides']['en']['body'].startswith('Loved')
        assert AuditLog.query.filter_by(action='notifications.template_override',
                                        entity_id=f'{key}:en').count() >= 1
        assert body(client.delete(f'/api/admin/notifications/templates/{key}', headers=auth(ops)))['data']['removed'] == 1
        assert _send(key, u.id, {'other_first': 'Amara'}).body == 'How was your trip with Amara?'
        assert client.put(f'/api/admin/notifications/templates/{key}', headers=auth(make_user('admin', admin_roles='finance')),
                          json={'lang': 'en', 'body': 'x'}).status_code == 403
    finally:
        O.query.filter_by(event_key=key).delete()
        db.session.commit()


# ── 6. receipts ─────────────────────────────────────────────────────────────

def test_tip_after_rating_issues_tip_receipt(client, auth, make_user):
    customer, driver, neg, rp = paid_carhire(make_user)
    receipt = RC.issue_and_send('carhire', neg.id)
    assert client.post(f'/api/rides/carhire/{neg.id}/tip', headers=auth(driver),
                       json={'amount_cents': 500}).status_code == 403
    assert body(client.post(f'/api/rides/carhire/{neg.id}/tip', headers=auth(customer),
                            json={'amount_cents': 10}))['data']['error_code'] == 'bad_amount_cents'
    r = client.post(f'/api/rides/carhire/{neg.id}/tip', headers=auth(customer), json={'amount_cents': 500})
    assert r.status_code == 201 and body(r)['data']['checkout_url'], r.get_json()
    tip_rp = db.session.get(RidePayment, body(r)['data']['ride_payment_id'])
    again = client.post(f'/api/rides/carhire/{neg.id}/tip', headers=auth(customer), json={'amount_cents': 500})
    assert again.status_code == 200 and body(again)['data']['reused'] is True
    email_provider.OUTBOX.clear()
    PS.handle_stripe_event(get_gateway().simulate_customer_pays(tip_rp.checkout_session_id))
    db.session.rollback()
    tr = TipReceipt.query.filter_by(ride_payment_id=tip_rp.id).one()
    assert tr.number.startswith(f'NR-TIP-{datetime.utcnow().year}-') and tr.amount_cents == 500
    assert tr.receipt_id == receipt.id and tr.emailed_at is not None
    mails = [m for m in email_provider.OUTBOX if m['to'] == customer.email and m['tag'] == 'tip.receipt']
    assert len(mails) == 1 and mails[0]['attachments'][0][0] == f'NegoRide-tip-{tr.number}.pdf'
    assert db.session.get(Negotiation, neg.id).tip_cents == 500
    payload = body(client.get(f'/api/rides/carhire/{neg.id}/receipt', headers=auth(customer)))['data']
    assert payload['tips']['total_cents'] == 500 and payload['tips']['items'][0]['number'] == tr.number
    assert payload['totals']['tip_cents'] == 0          # the issued receipt snapshot is never edited
    pdf = client.get(payload['tips']['items'][0]['pdf_url'], headers=auth(customer))
    assert pdf.status_code == 200 and pdf.data[:4] == b'%PDF'
    assert client.get(payload['tips']['items'][0]['pdf_url'], headers=auth(driver)).status_code == 403
    # idempotent
    assert RC.issue_tip_receipt(tip_rp.id).id == tr.id
    # tip window
    neg2 = paid_carhire(make_user, customer=customer, driver=driver)[2]
    n2 = db.session.get(Negotiation, neg2.id)
    n2.completed_at = datetime.utcnow() - timedelta(hours=80)
    db.session.commit()
    late = client.post(f'/api/rides/carhire/{neg2.id}/tip', headers=auth(customer), json={'amount_cents': 500})
    assert late.status_code == 410 and body(late)['data']['error_code'] == 'tip_window_closed'


def _pdf_text(blob):
    from pypdf import PdfReader
    # pypdf can insert a space after an uppercase glyph when reconstructing
    # positioned text from subset fonts (for example, "T est"). Remove only
    # that extraction artifact before checking receipt content/privacy.
    text = '\n'.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(blob)).pages)
    return re.sub(r'(?<=[A-Z])\s+(?=[a-z])', '', text)


def test_driver_receipt_pdf_is_redacted_and_email_has_logo_and_reply_to(client, auth, make_user, monkeypatch):
    monkeypatch.setenv('APP_URL', 'https://api.negoride.test')
    monkeypatch.delenv('EMAIL_LOGO_URL', raising=False)
    monkeypatch.delenv('EMAIL_REPLY_TO', raising=False)
    customer, driver, neg, rp = paid_carhire(make_user)
    email_provider.OUTBOX.clear()
    RC.issue_and_send('carhire', neg.id)
    mail = [m for m in email_provider.OUTBOX if m['to'] == customer.email][0]
    assert 'https://api.negoride.test/api/brand/logo.png' in mail['html'] and 'alt="NegoRide Canada"' in mail['html']
    assert 'v:roundrect' in mail['html'] and 'save-box' in mail['html']
    assert mail['reply_to'] == S.get('safety.support_email')
    full = _pdf_text(client.get(f'/api/rides/carhire/{neg.id}/receipt.pdf', headers=auth(customer)).data)
    drv = _pdf_text(client.get(f'/api/rides/carhire/{neg.id}/receipt.pdf', headers=auth(driver)).data)
    assert customer.last_name in full and '4242' in full
    assert 'DRIVER COPY' in drv and '4242' not in drv and f'{customer.first_name} {customer.last_name}' not in drv
    assert client.get('/api/brand/logo.png').status_code == 200
    assert RC.logo_data_uri().startswith('data:image/png;base64,')


def test_receipt_sweeper_and_admin_issue(client, auth, make_user):
    customer, driver, neg, rp = paid_carhire(make_user)
    r0 = db.session.get(RidePayment, rp.id)
    r0.captured_at = datetime.utcnow() - timedelta(minutes=5)
    db.session.commit()
    assert RC.find_receipt('carhire', neg.id) is None
    assert receipt_jobs.sweep_missing_receipts() >= 1
    db.session.rollback()
    assert RC.find_receipt('carhire', neg.id) is not None
    # a cancelled ride with a fee capture is not receipted by the sweeper
    c2, d2, neg2, rp2 = paid_carhire(make_user)
    db.session.get(Negotiation, neg2.id).trip_stage = 'CANCELLED_BY_CUSTOMER'
    x = db.session.get(RidePayment, rp2.id)
    x.captured_at = datetime.utcnow() - timedelta(minutes=5)
    db.session.commit()
    receipt_jobs.sweep_missing_receipts()
    db.session.rollback()
    assert RC.find_receipt('carhire', neg2.id) is None
    # admin "issue receipt" on the ride detail
    c3, d3, neg3, _ = paid_carhire(make_user)
    fin = make_user('admin', admin_roles='finance')
    r = client.post(f'/api/admin/rides/carhire/{neg3.id}/receipt/issue', headers=auth(fin))
    assert r.status_code == 200 and body(r)['data']['receipt']['number'].startswith('NR-'), r.get_json()
    assert AuditLog.query.filter_by(action='receipt.issue_requested', entity_id=str(neg3.id)).count() == 1
    r = client.post(f'/api/admin/rides/carhire/{neg3.id}/receipt/issue', headers=auth(fin))
    assert r.status_code == 200      # existing receipt returned, not re-issued
    assert Receipt.query.filter_by(ride_type='carhire', ride_id=neg3.id).count() == 1


def test_admin_resend_is_a_job(client, auth, make_user):
    customer, driver, neg, rp = paid_carhire(make_user)
    receipt = RC.issue_and_send('carhire', neg.id)
    sup = make_user('admin', admin_roles='support')
    r = client.post(f'/api/admin/receipts/{receipt.id}/resend', headers=auth(sup))
    assert r.status_code == 202 and body(r)['data']['queued'] is True and body(r)['data']['number'] == receipt.number
    assert AuditLog.query.filter_by(action='receipt.resend', entity_id=str(receipt.id)).count() == 1


def test_tax_rates_admin_is_effective_dated_and_audited(client, auth, make_user):
    fin = make_user('admin', admin_roles='finance')
    base = {'province': 'YT', 'name': 'Test future GST', 'gst_bp': 600, 'effective_from': '2099-01-01',
            'effective_to': '2099-06-30'}
    created = []
    # close the seeded open-ended YT row for the test (restored below)
    open_rows = TaxRate.query.filter(TaxRate.province == 'YT', TaxRate.effective_to.is_(None)).all()
    for t in open_rows:
        t.effective_to = date(2098, 12, 31)
    db.session.commit()
    try:
        r = client.post('/api/admin/finance/tax-rates', headers=auth(fin), json=base)
        assert r.status_code == 201, r.get_json()
        created.append(body(r)['data']['id'])
        assert body(r)['data']['total_bp'] == 600
        clash = client.post('/api/admin/finance/tax-rates', headers=auth(fin), json={**base, 'effective_from': '2099-03-01'})
        assert clash.status_code == 409 and body(clash)['data']['error_code'] == 'overlap'
        bad = client.post('/api/admin/finance/tax-rates', headers=auth(fin),
                          json={**base, 'hst_bp': 1300, 'effective_from': '2099-07-01', 'effective_to': None})
        assert body(bad)['data']['error_code'] == 'bad_tax_rate'
        r = client.post('/api/admin/finance/tax-rates', headers=auth(fin),
                        json={**base, 'gst_bp': 700, 'effective_from': '2099-07-01', 'effective_to': None})
        assert r.status_code == 201
        created.append(body(r)['data']['id'])
        r = client.put(f'/api/admin/finance/tax-rates/{created[0]}', headers=auth(fin), json={'gst_bp': 550})
        assert body(r)['data']['gst_bp'] == 550
        assert RC.rate_for('YT', date(2099, 2, 1)).gst_bp == 550
        assert RC.rate_for('YT', date(2099, 8, 1)).gst_bp == 700
        lst = body(client.get('/api/admin/finance/tax-rates?province=YT', headers=auth(fin)))['data']['items']
        assert {x['id'] for x in lst} >= set(created)
        assert AuditLog.query.filter_by(action='finance.tax_rate_update', entity_id=str(created[0])).count() == 1
        assert client.get('/api/admin/finance/tax-rates', headers=auth(make_user('customer'))).status_code == 403
    finally:
        db.session.rollback()
        TaxRate.query.filter(TaxRate.id.in_(created)).delete(synchronize_session=False)
        for t in TaxRate.query.filter(TaxRate.id.in_([r.id for r in open_rows])):
            t.effective_to = None
        db.session.commit()


def test_finance_exports_xlsx(client, auth, make_user):
    from openpyxl import load_workbook
    customer, driver, neg, rp = paid_carhire(make_user)
    RC.issue_and_send('carhire', neg.id)
    fin = make_user('admin', admin_roles='finance')
    for path in ('receipts', 'payments', 'refunds', 'commission', 'tax', 'payouts', 'credit-notes', 'statements'):
        r = client.get(f'/api/admin/finance/{path}?format=xlsx', headers=auth(fin))
        assert r.status_code == 200 and r.mimetype.endswith('spreadsheetml.sheet'), path
        wb = load_workbook(io.BytesIO(r.data))
        assert wb.active.max_row >= 1 and wb.active.cell(row=1, column=1).value, path
    csv_ = client.get('/api/admin/finance/receipts?format=csv', headers=auth(fin))
    assert csv_.mimetype == 'text/csv'
    assert AuditLog.query.filter(AuditLog.action == 'finance.export', AuditLog.actor_id == fin.id).count() >= 9


# ── 12. rating visibility re-evaluates the rules ──────────────────────────

def test_hiding_rating_reevaluates_rules(client, auth, make_user, monkeypatch):
    from backend.models.experience import RideRating
    from backend.services import account_service
    from tests.exp_helpers import make_carhire
    calls = []
    monkeypatch.setattr(account_service, 'evaluate_rating_rules', lambda did: calls.append(did))
    admin, d, c = make_user('admin'), make_user('driver'), make_user('customer')
    ride = make_carhire(c, d)
    client.post(f'/api/rides/carhire/{ride.id}/rating', headers=auth(c), json={'stars': 1})
    calls.clear()
    rr = RideRating.query.filter_by(ride_type='carhire', ride_id=ride.id).first()
    assert body(client.post(f'/api/admin/ratings/{rr.id}/hide', headers=auth(admin),
                            json={'reason': 'Discriminatory comment'}))['code'] == 1
    assert body(client.post(f'/api/admin/ratings/{rr.id}/unhide', headers=auth(admin),
                            json={'reason': 'Reviewed again, fine'}))['code'] == 1
    assert calls == [d.id, d.id]
